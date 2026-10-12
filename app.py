from __future__ import annotations

import ipaddress
import json
import re
import sqlite3
from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import streamlit as st

# Optional local file readers. The app still starts if an optional reader is
# missing; requirements.txt installs them for the complete experience.
try:
    from PIL import Image
except ImportError:  # pragma: no cover - optional dependency
    Image = None

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency
    pytesseract = None

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from docx import Document
except ImportError:  # pragma: no cover - optional dependency
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - optional dependency
    Presentation = None


# ============================================================
# AI SMARTGUARD — LOCAL, NO API KEY
# ============================================================

APP_DIR = Path(__file__).resolve().parent
DB_FILE = APP_DIR / "smartguard_history.db"
MAX_EXTRACTED_CHARS = 30_000

SUPPORTED_FILES = [
    "png", "jpg", "jpeg", "webp",
    "pdf", "docx", "txt", "csv", "md", "rtf",
    "pptx", "xlsx",
]

st.set_page_config(
    page_title="AI SmartGuard — Local",
    page_icon="🛡️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .stApp { background: #ffffff; }
    section[data-testid="stSidebar"] {
        background: #f7f7f8;
        border-right: 1px solid #e5e5e5;
    }
    section[data-testid="stSidebar"] > div { padding-top: 14px; }
    .brand {
        font-size: 19px; font-weight: 700; padding: 6px 8px 12px 8px;
        letter-spacing: -0.2px;
    }
    .brand-icon {
        display: inline-flex; width: 30px; height: 30px; align-items: center;
        justify-content: center; border-radius: 9px; background: #111827;
        color: white; margin-right: 8px;
    }
    .local-pill {
        display: inline-block; margin: 0 8px 12px; padding: 4px 8px;
        border: 1px solid #bbf7d0; border-radius: 999px; background: #f0fdf4;
        color: #166534; font-size: 11px; font-weight: 700; letter-spacing: .04em;
    }
    .sidebar-label {
        color: #6b7280; font-size: 12px; font-weight: 600;
        text-transform: uppercase; letter-spacing: .04em; margin: 18px 8px 7px;
    }
    .main-title {
        text-align: center; font-size: 32px; font-weight: 700;
        letter-spacing: -0.8px; margin-top: 10vh; margin-bottom: 8px;
        color: #171717;
    }
    .main-subtitle {
        text-align: center; color: #6b7280; font-size: 15px;
        margin: 0 auto 28px; max-width: 760px;
    }
    .composer-hint {
        text-align: center; color: #9ca3af; font-size: 12px; margin-top: 8px;
    }
    #MainMenu, footer { visibility: hidden; }
    header[data-testid="stHeader"] { background: transparent; }
    [data-testid="stChatMessage"] { padding-top: 12px; padding-bottom: 12px; }
    .stButton > button { border-radius: 10px; border: 1px solid #e5e7eb; }
    .block-container { padding-top: 1.5rem; }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# LOCAL HISTORY DATABASE
# ============================================================


def init_db() -> None:
    """Create or migrate the local history table without deleting old rows."""
    with sqlite3.connect(DB_FILE) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                analysis_type TEXT,
                content TEXT,
                classification TEXT,
                probability REAL,
                risk_level TEXT,
                summary TEXT
            )
            """
        )
        existing = {
            row[1] for row in conn.execute("PRAGMA table_info(history)").fetchall()
        }
        migrations = {
            "timestamp": "TEXT",
            "analysis_type": "TEXT",
            "content": "TEXT",
            "classification": "TEXT",
            "probability": "REAL",
            "risk_level": "TEXT",
            "summary": "TEXT",
        }
        for name, kind in migrations.items():
            if name not in existing:
                conn.execute(f"ALTER TABLE history ADD COLUMN {name} {kind}")
        conn.execute(
            "UPDATE history SET timestamp = ? WHERE timestamp IS NULL OR timestamp = ''",
            (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),),
        )
        conn.execute(
            "UPDATE history SET analysis_type = 'Text' WHERE analysis_type IS NULL"
        )
        conn.execute("UPDATE history SET summary = '' WHERE summary IS NULL")


def save_history(
    analysis_type: str,
    content: str,
    classification: str,
    probability: int,
    risk_level: str,
    summary: str,
) -> None:
    with sqlite3.connect(DB_FILE) as conn:
        conn.execute(
            """
            INSERT INTO history
                (timestamp, analysis_type, content, classification,
                 probability, risk_level, summary)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                analysis_type,
                content,
                classification,
                probability,
                risk_level,
                summary,
            ),
        )


def get_history(limit: int | None = None) -> list[tuple]:
    query = """
        SELECT id, timestamp, analysis_type, content, classification,
               probability, risk_level, summary
        FROM history
        ORDER BY id DESC
    """
    with sqlite3.connect(DB_FILE) as conn:
        if limit is None:
            return conn.execute(query).fetchall()
        return conn.execute(query + " LIMIT ?", (max(0, int(limit)),)).fetchall()


def get_stats() -> tuple[int, int, int, int, int, int]:
    with sqlite3.connect(DB_FILE) as conn:
        total = conn.execute("SELECT COUNT(*) FROM history").fetchone()[0]
        spam = conn.execute(
            "SELECT COUNT(*) FROM history WHERE classification = 'SPAM'"
        ).fetchone()[0]
        safe = conn.execute(
            "SELECT COUNT(*) FROM history WHERE classification = 'SAFE'"
        ).fetchone()[0]
        uncertain = conn.execute(
            "SELECT COUNT(*) FROM history WHERE classification = 'UNCERTAIN'"
        ).fetchone()[0]
        images = conn.execute(
            "SELECT COUNT(*) FROM history WHERE analysis_type = 'Image'"
        ).fetchone()[0]
        documents = conn.execute(
            "SELECT COUNT(*) FROM history WHERE analysis_type = 'Document'"
        ).fetchone()[0]
    return total, spam, safe, uncertain, images, documents


init_db()


# ============================================================
# LOCAL TEXT EXTRACTION
# ============================================================


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _clean_rtf(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\u-?\d+\??", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", text).strip()


def extract_upload_text(uploaded_file) -> tuple[str, str]:
    """Extract readable text locally; return (text, user-facing extraction note)."""
    if uploaded_file is None:
        return "", ""

    filename = uploaded_file.name or "uploaded file"
    suffix = Path(filename).suffix.lower()
    data = uploaded_file.getvalue()

    if suffix in {".txt", ".md", ".csv"}:
        return _decode_bytes(data)[:MAX_EXTRACTED_CHARS], ""

    if suffix == ".rtf":
        return _clean_rtf(_decode_bytes(data))[:MAX_EXTRACTED_CHARS], ""

    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        if Image is None or pytesseract is None:
            return "", (
                "Image OCR is unavailable. Install the Python packages and the "
                "Tesseract OCR application (see README.md) to read image text."
            )
        try:
            text = pytesseract.image_to_string(Image.open(BytesIO(data)))
            if text.strip():
                return text[:MAX_EXTRACTED_CHARS], "Text was read from the image using local OCR."
            return "", "No readable text was found in the image by local OCR."
        except Exception as exc:
            message = str(exc).lower()
            if "tesseract" in message and ("not installed" in message or "not found" in message):
                return "", (
                    "Tesseract OCR is not installed. Install it locally to scan "
                    "image text; this image was not sent anywhere."
                )
            return "", "Local OCR could not read this image. The image was not sent anywhere."

    if suffix == ".pdf":
        if PdfReader is None:
            return "", "PDF reader package is missing. Install the project requirements."
        try:
            reader = PdfReader(BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages[:80]]
            text = "\n".join(pages).strip()
            if text:
                return text[:MAX_EXTRACTED_CHARS], "PDF text was extracted locally."
            return "", "No selectable text was found in the PDF. Scanned PDFs need OCR."
        except Exception:
            return "", "This PDF could not be read locally (it may be encrypted or damaged)."

    if suffix == ".docx":
        if Document is None:
            return "", "DOCX reader package is missing. Install the project requirements."
        try:
            document = Document(BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts).strip()
            return text[:MAX_EXTRACTED_CHARS], "DOCX text was extracted locally."
        except Exception:
            return "", "This DOCX file could not be read locally."

    if suffix == ".xlsx":
        if load_workbook is None:
            return "", "XLSX reader package is missing. Install the project requirements."
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
            parts = []
            for sheet in workbook.worksheets:
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value is not None]
                    if values:
                        parts.append(" | ".join(values))
                    if sum(map(len, parts)) > MAX_EXTRACTED_CHARS:
                        break
                if sum(map(len, parts)) > MAX_EXTRACTED_CHARS:
                    break
            workbook.close()
            return "\n".join(parts)[:MAX_EXTRACTED_CHARS], "Spreadsheet text was extracted locally."
        except Exception:
            return "", "This spreadsheet could not be read locally."

    if suffix == ".pptx":
        if Presentation is None:
            return "", "PPTX reader package is missing. Install the project requirements."
        try:
            deck = Presentation(BytesIO(data))
            parts = []
            for slide_number, slide in enumerate(deck.slides, start=1):
                parts.append(f"Slide {slide_number}")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        parts.append(shape.text)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)[:MAX_EXTRACTED_CHARS], "Presentation text was extracted locally."
        except Exception:
            return "", "This presentation could not be read locally."

    return "", f"Text extraction is not configured for {suffix or 'this file type'}."


# ============================================================
# LOCAL SCAM RULES — NO NETWORK OR MODEL CALLS
# ============================================================

URL_PATTERN = re.compile(r"https?://[^\s<>\"']+|www\.[^\s<>\"']+", re.IGNORECASE)
SHORTENER_DOMAINS = {
    "bit.ly", "tinyurl.com", "t.co", "goo.gl", "is.gd", "cutt.ly",
    "rebrand.ly", "rb.gy", "tiny.cc", "ow.ly", "shorturl.at", "buff.ly",
}


def detect_urls(text: str) -> list[str]:
    urls = []
    for match in URL_PATTERN.findall(text or ""):
        candidate = match.rstrip(".,!?;:)]}>")
        if candidate and candidate not in urls:
            urls.append(candidate)
    return urls


def _get_host(url: str) -> str:
    parsed = urlparse(url if "://" in url else "https://" + url)
    return (parsed.hostname or "").lower().rstrip(".")


def inspect_links(urls: list[str]) -> tuple[str, list[str]]:
    if not urls:
        return "NONE", []

    reasons: list[str] = []
    risk_score = 0
    for url in urls:
        host = _get_host(url)
        if not host:
            continue
        try:
            ipaddress.ip_address(host)
            if "A URL uses a raw IP address instead of a recognizable domain." not in reasons:
                reasons.append("A URL uses a raw IP address instead of a recognizable domain.")
            risk_score = max(risk_score, 3)
        except ValueError:
            pass

        if host.startswith("xn--") or ".xn--" in host:
            if "A URL contains an internationalized (punycode) domain that should be checked carefully." not in reasons:
                reasons.append("A URL contains an internationalized (punycode) domain that should be checked carefully.")
            risk_score = max(risk_score, 3)

        if host in SHORTENER_DOMAINS:
            if "A link-shortening service hides the final destination." not in reasons:
                reasons.append("A link-shortening service hides the final destination.")
            risk_score = max(risk_score, 2)

    if risk_score >= 3:
        return "SUSPICIOUS", reasons
    if risk_score >= 2:
        return "LOW", reasons
    return "LOW", ["A URL is present. Its destination was not opened or verified."]


def _has_unnegated_match(pattern: str, text: str) -> bool:
    """Match a request/action, but ignore common safety warnings such as 'never share'."""
    for match in re.finditer(pattern, text, flags=re.IGNORECASE):
        prefix = text[max(0, match.start() - 45):match.start()]
        if re.search(
            r"\b(?:never|don't|dont|do not|must not|should not|cannot|can't|will not|won't|avoid)"
            r"(?:\s+\w+){0,4}\s*$",
            prefix,
            flags=re.IGNORECASE,
        ):
            continue
        if re.search(r"\b(?:warn|remind|advise)\w*\s+(?:you\s+)?(?:to\s+)?(?:never|not)\s*$", prefix, re.IGNORECASE):
            continue
        return True
    return False


def analyze_locally(text: str, extraction_note: str = "") -> dict:
    """Apply transparent, conservative rules to supplied text; no data leaves the app."""
    clean_text = (text or "").strip()
    flat = re.sub(r"\s+", " ", clean_text.casefold())
    urls = detect_urls(clean_text)
    link_risk, link_reasons = inspect_links(urls)

    if not clean_text:
        return {
            "classification": "UNCERTAIN",
            "scam_probability": 50,
            "risk_level": "UNKNOWN",
            "summary": "There was no readable text to scan, so this result is inconclusive.",
            "reasons": [extraction_note or "No text was provided or extracted."],
            "link_analysis": {
                "links_detected": bool(urls), "risk": link_risk,
                "reasons": link_reasons, "urls": urls,
            },
            "recommendation": "Do not treat this as a safe result. Try a clearer image or a text-based file.",
            "extracted_text": "",
            "extraction_note": extraction_note,
        }

    score = 0
    reasons: list[str] = []

    def add_signal(weight: int, reason: str) -> None:
        nonlocal score
        if reason not in reasons:
            score += weight
            reasons.append(reason)

    secret_request = (
        r"\b(?:share|send|tell|provide|reveal|forward|reply\s+with|text|message|give)\b"
        r".{0,45}\b(?:otp|one[\s-]?time password|verification code|security code|"
        r"upi pin|pin|password|passcode|cvv|card number)\b"
    )
    if _has_unnegated_match(secret_request, flat):
        add_signal(7, "The message appears to ask someone to disclose a secret such as an OTP, PIN, password, or card security detail.")

    if re.search(
        r"\b(?:enter|share|provide|tell|type)\b.{0,35}\bupi\s*pin\b.{0,45}\b(?:receive|get|claim|accept)\b"
        r"|\b(?:receive|get|claim)\b.{0,45}\bupi\s*pin\b.{0,35}\b(?:enter|share|provide|tell)\b",
        flat,
    ):
        add_signal(7, "The message links a UPI PIN to receiving money; a UPI PIN should not be needed to receive funds.")

    if re.search(r"\b(?:anydesk|teamviewer|quick.?support|remote access|screen.?sharing app)\b", flat) and re.search(
        r"\b(?:install|download|open|allow|connect|share your screen)\b", flat
    ):
        add_signal(6, "The message asks for remote-access or screen-sharing software, which can expose accounts and payments.")

    account_lock = re.search(
        r"\b(?:account|bank account|card|wallet|upi|kyc)\b.{0,55}\b(?:block(?:ed)?|suspend(?:ed)?|freeze|frozen|expire[sd]?|deactivat(?:e|ed))\b",
        flat,
    ) or re.search(
        r"\b(?:block(?:ed)?|suspend(?:ed)?|freeze|frozen|expire[sd]?|deactivat(?:e|ed))\b.{0,55}\b(?:account|bank|card|wallet|upi|kyc)\b",
        flat,
    )
    account_action = re.search(r"\b(?:click|tap|verify|update|reactivate|unlock|confirm|link|otp|password|pin)\b", flat)
    if account_lock and (account_action or urls):
        add_signal(5, "An account-blocking or KYC threat is paired with a request to act, verify, or follow a link.")

    verification_terms = re.search(
        r"\b(?:verify|update|reactivate|unlock|confirm|validate|restore)\b.{0,45}"
        r"\b(?:account|kyc|bank|wallet|identity|card|details|payment)\b"
        r"|\b(?:account|kyc|bank|wallet|identity|card|details|payment)\b.{0,45}"
        r"\b(?:verify|update|reactivate|unlock|confirm|validate|restore)\b",
        flat,
    )
    if urls and verification_terms:
        add_signal(3, "A link is combined with account, identity, or payment verification language.")

    prize = re.search(r"\b(?:won|winner|prize|lottery|reward|gift voucher|free gift|cashback)\b", flat)
    if prize and (urls or re.search(r"\b(?:fee|pay|payment|claim|limited time|selected)\b", flat)):
        add_signal(5, "A prize, reward, or giveaway is paired with a claim link, payment, or urgency.")

    job_offer = re.search(r"\b(?:job offer|work from home|remote job|selected for the job|guaranteed income)\b", flat)
    upfront_fee = re.search(r"\b(?:registration fee|training fee|security deposit|processing fee|pay upfront|advance fee)\b", flat)
    if job_offer and upfront_fee:
        add_signal(6, "A job offer is paired with an upfront fee or deposit, a common job-scam pattern.")

    loan_offer = re.search(r"\b(?:instant loan|loan approved|pre.?approved loan|loan offer)\b", flat)
    if loan_offer and re.search(r"\b(?:advance fee|processing fee|pay first|upfront|insurance fee)\b", flat):
        add_signal(5, "A loan offer is paired with a fee demanded before funds are received.")

    refund = re.search(r"\b(?:refund|cashback|reimbursement)\b", flat)
    if refund and re.search(r"\b(?:otp|upi pin|cvv|card number|pay a fee|processing fee|click here)\b", flat):
        add_signal(5, "A refund is linked to a request for sensitive payment details, a fee, or a risky action.")

    if re.search(r"\b(?:guaranteed returns?|double your money|risk.?free profit|100% profit)\b", flat):
        add_signal(4, "The message promises unusually high or guaranteed investment returns.")

    urgency = re.search(
        r"\b(?:within\s+\d+\s*(?:minutes?|hours?)|act now|immediately|urgent|last warning|final notice)\b",
        flat,
    )
    if urgency and (account_lock or urls or secret_request):
        add_signal(2, "Urgency or a short deadline is used alongside a sensitive request or suspicious action.")

    if urls:
        for link_reason in link_reasons:
            if "link-shortening" in link_reason.lower():
                add_signal(2, link_reason)
            elif "raw ip" in link_reason.lower() or "punycode" in link_reason.lower():
                add_signal(3, link_reason)

    if score >= 10:
        classification = "SPAM"
        risk_level = "CRITICAL"
        probability = min(98, 92 + min(score - 10, 6))
        recommendation = "Do not click, reply, pay, or share codes or account details. Contact the organization using its official app or a number you already trust."
        summary = "Several strong scam indicators were found in this text."
    elif score >= 6:
        classification = "SPAM"
        risk_level = "HIGH"
        probability = min(91, 72 + (score - 6) * 4)
        recommendation = "Do not use the link or contact details in the message. Verify independently through the organization’s official app or website."
        summary = "The text contains one or more strong scam indicators."
    elif score >= 3:
        classification = "UNCERTAIN"
        risk_level = "MEDIUM"
        probability = min(65, 44 + (score - 3) * 6)
        recommendation = "Treat it cautiously. Avoid links and requests for money or sensitive information until you verify through an official channel."
        summary = "Some warning signs were found, but the local rules cannot confidently classify this message."
    else:
        classification = "SAFE"
        risk_level = "LOW"
        probability = 8 + score * 10
        recommendation = "No strong scam indicators were detected. This is not proof that the message is genuine; verify unexpected requests independently."
        summary = "No strong scam indicators were found by the local rules."
        reasons = ["No strong scam indicators detected in the provided text."]

    return {
        "classification": classification,
        "scam_probability": int(probability),
        "risk_level": risk_level,
        "summary": summary,
        "reasons": reasons,
        "link_analysis": {
            "links_detected": bool(urls),
            "risk": link_risk,
            "reasons": link_reasons,
            "urls": urls,
        },
        "recommendation": recommendation,
        "extracted_text": clean_text[:6000],
        "extraction_note": extraction_note,
    }


def analyze_submission(text: str, uploaded_file=None) -> tuple[dict, str]:
    extracted, extraction_note = extract_upload_text(uploaded_file)
    parts = [part for part in (text.strip(), extracted.strip()) if part]
    scan_text = "\n\n".join(parts)
    result = analyze_locally(scan_text, extraction_note=extraction_note)
    return result, extracted


# ============================================================
# RESULT DISPLAY
# ============================================================


def render_result(result: dict) -> None:
    classification = str(result.get("classification", "UNCERTAIN")).upper()
    probability = result.get("scam_probability", 50)
    risk = result.get("risk_level", "UNKNOWN")

    if classification == "SPAM":
        st.error("🚨 SPAM / SCAM INDICATORS FOUND")
    elif classification == "SAFE":
        st.success("✅ NO STRONG SCAM INDICATORS FOUND")
    else:
        st.warning("⚠️ UNCERTAIN — REVIEW CAREFULLY")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("Local result", classification)
    with c2:
        st.metric("Heuristic estimate", f"{probability}%")
    with c3:
        st.metric("Risk level", risk)

    st.caption("The percentage is a rule-based estimate, not a calibrated AI probability.")
    st.markdown("#### 🧠 Summary")
    st.write(result.get("summary", ""))

    reasons = result.get("reasons", [])
    if reasons:
        st.markdown("#### ⚠️ Detection reasons")
        for reason in reasons:
            st.write(f"• {reason}")

    links = result.get("link_analysis", {})
    st.markdown("#### 🔗 Link analysis")
    if links.get("links_detected"):
        st.write(f"**Link risk:** {links.get('risk', 'UNKNOWN')} — URLs were not opened.")
        for url in links.get("urls", []):
            st.code(url, language=None)
        for reason in links.get("reasons", []):
            st.write(f"• {reason}")
    else:
        st.write("No URL detected.")

    st.markdown("#### 🛡️ Recommendation")
    st.info(result.get("recommendation", "Verify unexpected requests independently."))

    note = result.get("extraction_note")
    if note:
        st.caption(note)
    extracted = result.get("extracted_text", "")
    if extracted:
        with st.expander("Text scanned locally"):
            st.code(extracted, language=None)


# ============================================================
# SESSION STATE
# ============================================================

if "messages" not in st.session_state:
    st.session_state.messages = []
if "chat_title" not in st.session_state:
    st.session_state.chat_title = "New detection"
if "show_history" not in st.session_state:
    st.session_state.show_history = False


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.markdown(
        '<div class="brand"><span class="brand-icon">🛡️</span>AI SmartGuard</div>',
        unsafe_allow_html=True,
    )
    st.markdown('<div class="local-pill">LOCAL · NO API KEY</div>', unsafe_allow_html=True)

    if st.button("＋  New detection", use_container_width=True):
        st.session_state.messages = []
        st.session_state.chat_title = "New detection"
        st.session_state.show_history = False
        st.rerun()

    st.markdown('<div class="sidebar-label">Workspace</div>', unsafe_allow_html=True)
    if st.button("📜 History", use_container_width=True):
        st.session_state.show_history = True
    if st.button("💬 Current chat", use_container_width=True):
        st.session_state.show_history = False

    st.markdown('<div class="sidebar-label">Recent history</div>', unsafe_allow_html=True)
    recent = get_history(limit=8)
    if not recent:
        st.caption("No detections yet.")
    else:
        for row in recent:
            _, timestamp, kind, content, classification, _, _, _ = row
            icon = "🚨" if classification == "SPAM" else ("✅" if classification == "SAFE" else "⚠️")
            label = (content or "File analysis").strip().replace("\n", " ")
            if len(label) > 34:
                label = label[:34] + "…"
            st.caption(f"{icon} {label}")
            st.caption(f"{kind} · {timestamp}")

    st.markdown('<div class="sidebar-label">Your usage</div>', unsafe_allow_html=True)
    total, spam_count, safe_count, uncertain_count, image_count, document_count = get_stats()
    c1, c2 = st.columns(2)
    with c1:
        st.metric("Analyses", total)
    with c2:
        st.metric("Scams", spam_count)
    st.caption(
        f"Safe: {safe_count} · Review: {uncertain_count} · "
        f"Images: {image_count} · Documents: {document_count}"
    )
    st.divider()
    st.caption("Rule-based scanning runs on this machine. Text and uploads are not sent to an AI service.")


# ============================================================
# HISTORY PAGE
# ============================================================

if st.session_state.show_history:
    st.title("History")
    st.caption("Analysis records are kept in a local SQLite file next to app.py.")

    total, spam_count, safe_count, uncertain_count, image_count, document_count = get_stats()
    cols = st.columns(6)
    metrics = [
        ("Total", total), ("🚨 Spam", spam_count), ("✅ Safe", safe_count),
        ("⚠️ Review", uncertain_count), ("🖼️ Images", image_count),
        ("📄 Documents", document_count),
    ]
    for col, (label, value) in zip(cols, metrics):
        with col:
            st.metric(label, value)

    st.divider()
    rows = get_history()
    if not rows:
        st.info("No history yet.")
    else:
        for row in rows:
            row_id, timestamp, kind, content, classification, probability, risk, summary = row
            icon = "🚨" if classification == "SPAM" else ("✅" if classification == "SAFE" else "⚠️")
            with st.expander(f"{icon} {classification} · {kind} · {timestamp}"):
                st.write(f"**Heuristic estimate:** {probability}%")
                st.write(f"**Risk level:** {risk}")
                st.write(f"**Summary:** {summary}")
                st.write("**Submitted text or filename:**")
                st.code(content or "File analysis", language=None)

    st.stop()


# ============================================================
# MAIN CHAT
# ============================================================

if not st.session_state.messages:
    st.markdown('<div class="main-title">How can I help protect you?</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="main-subtitle">Paste a suspicious message or attach a photo, '
        'PDF, document, spreadsheet, or presentation. SmartGuard checks the content '
        'locally using transparent scam indicators.</div>',
        unsafe_allow_html=True,
    )
else:
    st.markdown(f"### {st.session_state.chat_title}")

for message in st.session_state.messages:
    role = message["role"]
    with st.chat_message(role, avatar="🧑" if role == "user" else "🛡️"):
        if message.get("image") is not None:
            st.image(message["image"], width=360)
        if message.get("file_name") and message.get("image") is None:
            st.caption(f"📄 {message['file_name']}")
        if message.get("text"):
            st.markdown(message["text"])
        if role == "assistant" and message.get("result"):
            render_result(message["result"])

prompt = st.chat_input(
    "Message AI SmartGuard…",
    accept_file=True,
    file_type=SUPPORTED_FILES,
)

if prompt:
    text = getattr(prompt, "text", "") or ""
    files = getattr(prompt, "files", []) or []
    uploaded_file = files[0] if files else None

    if len(files) > 1:
        st.warning("This version analyzes one attachment per message. The first file will be scanned.")
    if not text.strip() and uploaded_file is None:
        st.warning("Enter a message or attach a photo or file.")
        st.stop()

    is_image = uploaded_file is not None and (uploaded_file.type or "").startswith("image/")
    is_document = uploaded_file is not None and not is_image
    image_bytes = uploaded_file.getvalue() if is_image else None

    with st.chat_message("user", avatar="🧑"):
        if uploaded_file is not None:
            if is_image:
                st.image(image_bytes, width=360)
            else:
                st.caption(f"📄 {uploaded_file.name}")
        if text.strip():
            st.markdown(text)

    st.session_state.messages.append(
        {
            "role": "user",
            "text": text,
            "image": image_bytes,
            "file_name": uploaded_file.name if uploaded_file else None,
        }
    )

    with st.chat_message("assistant", avatar="🛡️"):
        with st.spinner("Scanning locally…"):
            try:
                result, extracted_text = analyze_submission(text, uploaded_file)
                if uploaded_file is not None:
                    analysis_type = "Image" if is_image else "Document"
                else:
                    analysis_type = "Text"
                history_content = text.strip() or (uploaded_file.name if uploaded_file else "Analysis")
                save_history(
                    analysis_type=analysis_type,
                    content=history_content,
                    classification=result["classification"],
                    probability=result["scam_probability"],
                    risk_level=result["risk_level"],
                    summary=result["summary"],
                )
                render_result(result)
                st.session_state.messages.append(
                    {
                        "role": "assistant",
                        "text": "",
                        "image": None,
                        "result": result,
                    }
                )
                if st.session_state.chat_title == "New detection":
                    title = text.strip().replace("\n", " ") if text.strip() else (uploaded_file.name if uploaded_file else "Analysis")
                    st.session_state.chat_title = title[:40] + ("…" if len(title) > 40 else "")
            except Exception as exc:
                st.error("The local scan could not be completed.")
                st.caption(str(exc))

st.markdown(
    '<div class="composer-hint">Local rules can miss new scams or flag unusual legitimate messages. '
    'Never share OTPs, passwords, PINs, or other sensitive information.</div>',
    unsafe_allow_html=True,
)
